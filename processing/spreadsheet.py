"""Read spreadsheet evidence. Never execute formulas, macros or external links."""
import json
import sys
import zipfile

MAX_CELLS = 200_000
MAX_TEXT = 8_000_000


def extract(path, kind):
    pages = []
    cells = 0
    length = 0

    def page(name, entries):
        nonlocal length
        text = "## 工作表：" + name + "\n" + "\n".join(entries)
        length += len(text.encode("utf-8"))
        if length > MAX_TEXT:
            raise ValueError("content_limit")
        if entries:
            pages.append({"text": text, "page": None})

    if kind == "xlsx":
        import openpyxl
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 20_000 or sum(i.file_size for i in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("content_limit")
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if len(book.worksheets) > 200:
                raise ValueError("content_limit")
            for sheet, values in zip(book.worksheets, cached.worksheets):
                if sheet.sheet_state != "visible":
                    continue
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                    raise ValueError("content_limit")
                entries = []
                for row, cached_row in zip(sheet.iter_rows(), values.iter_rows()):
                    for cell, value in zip(row, cached_row):
                        cells += 1
                        if cells > MAX_CELLS:
                            raise ValueError("content_limit")
                        if cell.value is None:
                            continue
                        text = str(getattr(cell.value, "text", cell.value))
                        if cell.data_type == "f":
                            text = f"公式源码（未执行）：`{text}`；文件缓存值：{value.value if value.value is not None else '无'}"
                        entries.append(f"{cell.coordinate}：{text}")
                page(sheet.title, entries)
        finally:
            book.close()
            cached.close()
    else:
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            if book.nsheets > 200:
                raise ValueError("content_limit")
            for sheet in book.sheets():
                if sheet.visibility:
                    continue
                if sheet.nrows * sheet.ncols + cells > MAX_CELLS:
                    raise ValueError("content_limit")
                entries = []
                for r in range(sheet.nrows):
                    for c in range(sheet.ncols):
                        cells += 1
                        value = sheet.cell_value(r, c)
                        if value != "":
                            entries.append(f"第{r+1}行第{c+1}列：{value}")
                page(sheet.name, entries)
        finally:
            book.release_resources()
    return pages


if __name__ == "__main__":
    try:
        result = {"pages": extract(sys.argv[1], sys.argv[2])}
    except ImportError:
        result = {"error": "missing_dependency"}
    except Exception as error:
        result = {"error": "content_limit" if str(error) == "content_limit" else "invalid_file"}
    print(json.dumps(result, ensure_ascii=True))
